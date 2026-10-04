"""results_table.py — Lưu kết quả JSON và tự động xuất ra experiments.xlsx từ mẫu.

Sheet "Experiments":
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
"""
from __future__ import annotations

import json
from pathlib import Path
import openpyxl

FORMULA_COLUMNS = {
    "step0_gap_vs_lnC",
    "gap_val_minus_train",
    "delta_val_f1_vs_base",
    "beyond_noise",
}


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (không ghi best_state) ra JSON."""
    p = Path(results_dir)
    p.mkdir(parents=True, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    out_file = p / f"{exp_id}.json"
    data_to_save = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"],
    }
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2, ensure_ascii=False)
    return str(out_file)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict sắp theo exp_id."""
    p = Path(results_dir)
    if not p.exists():
        return []
    res = []
    for f in sorted(p.glob("*.json")):
        with open(f, "r", encoding="utf-8") as fp:
            res.append(json.load(fp))
    return sorted(res, key=lambda x: x["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary."""
    cfg = result["cfg"]
    summary = result.get("summary", {})
    exp_id = cfg["exp_id"]
    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce"),
        "optimizer": cfg.get("optimizer", ""),
        "lr": cfg.get("lr", ""),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": str(cfg.get("hidden", (256, 128))),
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": str(cfg.get("clip_norm", "None")),
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": summary.get("step0_loss", ""),
        "best_val_loss": summary.get("best_val_loss", ""),
        "best_epoch": summary.get("best_epoch", ""),
        "final_train_loss": summary.get("final_train_loss", ""),
        "final_val_loss": summary.get("final_val_loss", ""),
        "val_acc": summary.get("val_acc", ""),
        "val_macro_f1": summary.get("val_macro_f1", ""),
        "time_per_epoch_s": summary.get("time_per_epoch_s", ""),
        "peak_mem_MB": summary.get("peak_mem_MB", 0.0),
        "diverged": summary.get("diverged", False),
        "eval_acc": eval_scores.get("acc", "") if eval_scores else "",
        "eval_macro_f1": eval_scores.get("macro_f1", "") if eval_scores else "",
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet 'Experiments' của mẫu mà không ghi đè cột công thức."""
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    col_map = {}
    for col_idx in range(1, ws.max_column + 1):
        header_val = ws.cell(row=1, column=col_idx).value
        if header_val:
            col_map[str(header_val).strip()] = col_idx

    for row_num, row_data in enumerate(rows, start=2):
        for key, val in row_data.items():
            if key in col_map and key not in FORMULA_COLUMNS:
                col_idx = col_map[key]
                ws.cell(row=row_num, column=col_idx, value=val)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
